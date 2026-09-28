#!/usr/bin/env python3
"""
check_datasets.py

Final step of run_all.sh. Checks every Supplementary Dataset workbook in
datasets/ so that the released files open cleanly in Excel and in other
readers:

  1. Removes package relationships that point to parts missing from the
     workbook (for example a drawing or comment part), with the matching
     <drawing>/<legacyDrawing> elements and [Content_Types].xml overrides.
     Such dangling links make Excel offer a repair and stop openpyxl.
  2. Opens each workbook with openpyxl and fails if it cannot be read.
  3. Fails if any cell holds a comment, or if a cell contains a term that is
     not used in the paper (see RETIRED_TERMS).

The data themselves are not changed.
"""

import glob
import os
import posixpath
import re
import shutil
import sys
import tempfile
import zipfile

import openpyxl

DS = "datasets"
# wording replaced in the paper; a Dataset that still uses it is out of date
RETIRED_TERMS = ["network of record", "networks of record", "instances"]

REL_RE = re.compile(r'<Relationship\b[^>]*?/>')
ATTR = lambda name, s: (re.search(r'\b%s="([^"]*)"' % name, s) or [None, None])[1]


def resolve(rels_path, target):
    """Part name that a relationship target in rels_path points to."""
    if target.startswith("/"):
        return target.lstrip("/")
    src_dir = posixpath.dirname(posixpath.dirname(rels_path))  # folder of the source part
    return posixpath.normpath(posixpath.join(src_dir, target))


def sanitize(path):
    with zipfile.ZipFile(path) as z:
        parts = {i.filename: z.read(i.filename) for i in z.infolist()}
    names = set(parts)
    removed = []
    for rp in [p for p in names if p.endswith(".rels")]:
        xml = parts[rp].decode("utf-8")
        drop_ids = []
        for rel in REL_RE.findall(xml):
            if ATTR("TargetMode", rel) == "External":
                continue
            if resolve(rp, ATTR("Target", rel)) not in names:
                drop_ids.append(ATTR("Id", rel))
                xml = xml.replace(rel, "")
        if not drop_ids:
            continue
        parts[rp] = xml.encode("utf-8")
        removed += ["%s:%s" % (rp, i) for i in drop_ids]
        src = posixpath.join(posixpath.dirname(posixpath.dirname(rp)),
                             posixpath.basename(rp)[:-len(".rels")])
        if src in parts:
            sx = parts[src].decode("utf-8")
            for i in drop_ids:
                sx = re.sub(r'<(drawing|legacyDrawing|legacyDrawingHF)\b[^>]*r:id="%s"[^>]*/>'
                            % re.escape(i), "", sx)
            parts[src] = sx.encode("utf-8")
    ct = parts["[Content_Types].xml"].decode("utf-8")
    for ov in re.findall(r'<Override\b[^>]*/>', ct):
        if ATTR("PartName", ov).lstrip("/") not in names:
            ct = ct.replace(ov, "")
            removed.append("override " + ATTR("PartName", ov))
    parts["[Content_Types].xml"] = ct.encode("utf-8")
    if removed:
        fd, tmp = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
            order = ["[Content_Types].xml"] + [p for p in parts if p != "[Content_Types].xml"]
            for p in order:
                z.writestr(p, parts[p])
        shutil.move(tmp, path)
    return removed


def check(path):
    problems = []
    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.comment is not None:
                    problems.append("%s!%s has a cell comment" % (ws.title, c.coordinate))
                if isinstance(c.value, str):
                    low = c.value.lower()
                    for t in RETIRED_TERMS:
                        if re.search(r"\b%s\b" % re.escape(t), low):
                            problems.append("%s!%s uses '%s'" % (ws.title, c.coordinate, t))
        if any(t in ws.title.lower() for t in RETIRED_TERMS):
            problems.append("sheet name '%s'" % ws.title)
    return wb, problems


def main():
    files = sorted(glob.glob(os.path.join(DS, "Dataset_S*.xlsx")))
    if len(files) != 8:
        sys.exit("STOP: expected 8 Datasets in %s, found %d" % (DS, len(files)))
    bad = 0
    for f in files:
        removed = sanitize(f)
        try:
            wb, problems = check(f)
        except Exception as e:  # unreadable workbook
            print("FAIL %s: %s" % (os.path.basename(f), e))
            bad += 1
            continue
        status = "ok" if not problems else "FAIL"
        print("%-4s %-50s %2d sheets%s" % (status, os.path.basename(f), len(wb.sheetnames),
              ", removed %d dangling links" % len(removed) if removed else ""))
        for p in problems[:20]:
            print("       " + p)
        bad += bool(problems)
    if bad:
        sys.exit("STOP: %d Dataset(s) failed the check" % bad)
    print("all Datasets open cleanly")


if __name__ == "__main__":
    main()
